# -*- coding: utf-8 -*-
"""
Formato de lectura de las hojas de salida (ancho, negrita, panel fijo).

No toca ni un valor: solo el aspecto con el que se abren los dos
libros. El calculo ya esta hecho cuando esto corre -- si alguna vez
alguna de estas funciones cambia un dato, es un bug, no una opcion.

Pedido del usuario: "necesito que ordenes las planillas". Antes las
hojas salian tal cual las deja pandas: encabezados sin negrita,
columnas de ancho 8 (todos los numeros como ####), sin panel fijo, asi
que revisar 100.000 filas obligaba a formatear a mano cada vez que se
regeneraba el libro.
"""

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


# Ancho de columna en "caracteres" de Excel. El maximo existe para que
# una Barra o un nombre de central largo no empuje el resto de la hoja
# fuera de la pantalla.
ANCHO_MINIMO = 9
ANCHO_MAXIMO = 38

# Cuantas filas de datos se miran para decidir ancho y formato de
# numero. Con 200 alcanza para ver la forma de la columna y no cuesta
# nada en una hoja de 100.000 filas.
FILAS_MUESTRA = 200

# Tope de celdas a las que se les escribe el formato de numero una por
# una (openpyxl no tiene formato por columna que Excel respete). El
# tope esta puesto MUY arriba a proposito: un mes 15-minutal de una
# decena de centrales son ~30.000 filas, o sea ~1,5 millones de celdas
# en "Calculo E Costos" (medido: ~1 segundo cada millon y medio). Es
# un freno para un caso disparatado, no algo que un periodo normal
# tenga que tocar -- si lo toca, el libro sale igual, solo que sin
# separador de miles.
MAXIMO_CELDAS_FORMATO = 8_000_000

FORMATO_ENTERO = "#,##0"
FORMATO_DECIMAL = "#,##0.00"
FORMATO_FECHA = "yyyy-mm-dd hh:mm"


def _es_numero(valor):
    return isinstance(valor, (int, float)) and not isinstance(valor, bool)


def _formato_de_columna(valores):
    """
    Decide el formato de numero de una columna a partir de una
    muestra de sus valores. Devuelve None si no hay que tocarla
    (columna de texto, vacia, o mezcla de tipos).
    """

    con_valor = [v for v in valores if v is not None and v != ""]

    if not con_valor:
        return None

    if all(hasattr(v, "year") and hasattr(v, "month") for v in con_valor):
        return FORMATO_FECHA

    if not all(_es_numero(v) for v in con_valor):
        return None

    if any(float(v) != int(v) for v in con_valor):
        return FORMATO_DECIMAL

    return FORMATO_ENTERO


def _texto_de_celda(valor):
    if valor is None:
        return ""
    if isinstance(valor, float):
        # Un 1234.5678901 mide 12 caracteres en memoria pero se ve
        # como 1.234,57 una vez aplicado FORMATO_DECIMAL.
        return f"{valor:,.2f}"
    return str(valor)


def formatear_hoja(ws, fila_encabezado=1):
    """
    Deja una hoja lista para leer:

      - las filas 1..fila_encabezado (titulos de cuadro, encabezados
        de grupo y nombres de columna) en negrita y centradas;
      - panel inmovilizado justo debajo del encabezado, para que los
        nombres de columna sigan a la vista al bajar;
      - ancho de columna segun lo que realmente hay adentro;
      - separador de miles (y dos decimales cuando la columna los
        tiene) en las columnas numericas, y fecha legible en las de
        fecha.

    fila_encabezado: fila (1-indexada) donde estan los NOMBRES de
    columna. Es 1 en las hojas que escribe pandas solas, 2 en las que
    llevan una fila de encabezados de grupo arriba (Calculo E Costos,
    Calculo RE545, Ofertas SSCC) y 3 en PRORRATA_RETIROS (titulo de
    la hoja + titulo de cada cuadro).
    """

    if ws.max_row <= fila_encabezado:
        # Hoja vacia o sin filas de datos: no hay nada que ordenar.
        return

    for fila in ws.iter_rows(min_row=1, max_row=fila_encabezado):
        for celda in fila:
            if celda.value is None:
                # Tambien son las celdas "de adentro" de un merge, que
                # no aceptan estilo.
                continue
            celda.font = Font(bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=1)

    primera_fila_datos = fila_encabezado + 1
    ultima_muestra = min(ws.max_row, primera_fila_datos + FILAS_MUESTRA - 1)

    muestra = [
        [celda.value for celda in fila]
        for fila in ws.iter_rows(
            min_row=primera_fila_datos, max_row=ultima_muestra
        )
    ]

    formatos = {}

    for indice in range(ws.max_column):

        columna = [
            fila[indice] for fila in muestra
            if indice < len(fila)
        ]

        titulos = [
            ws.cell(row=fila, column=indice + 1).value
            for fila in range(1, fila_encabezado + 1)
        ]

        anchos = [
            len(_texto_de_celda(valor))
            for valor in columna + titulos
        ]

        letra = get_column_letter(indice + 1)
        ws.column_dimensions[letra].width = min(
            ANCHO_MAXIMO, max(ANCHO_MINIMO, max(anchos, default=0) + 2)
        )

        formato = _formato_de_columna(columna)
        if formato:
            formatos[indice + 1] = formato

    if not formatos:
        return

    celdas = (ws.max_row - fila_encabezado) * len(formatos)
    if celdas > MAXIMO_CELDAS_FORMATO:
        return

    for fila in ws.iter_rows(
        min_row=primera_fila_datos,
        min_col=min(formatos),
        max_col=max(formatos),
    ):
        for celda in fila:
            formato = formatos.get(celda.column)
            if formato and celda.value is not None:
                celda.number_format = formato


def formatear_libro(wb, filas_encabezado=None):
    """
    Aplica formatear_hoja() a todas las hojas del libro.

    filas_encabezado: {nombre de hoja: fila de los nombres de
    columna}. Las hojas que no esten en el diccionario usan 1.
    """

    filas_encabezado = filas_encabezado or {}

    for nombre in wb.sheetnames:
        formatear_hoja(wb[nombre], filas_encabezado.get(nombre, 1))
