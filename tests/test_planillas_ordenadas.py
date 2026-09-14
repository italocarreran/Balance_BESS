# -*- coding: utf-8 -*-
"""
Las planillas ordenadas: que hojas salen, en que orden, con que
columnas y con que formato -- y que nada de eso toca un numero.

Pedido del usuario: "sin modificar el calculo, necesito que ordenes
las planillas, quita los auxiliares innecesarios".
"""

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import Script.nucleo as nucleo  # noqa: E402


class TestMedidoresSinAuxiliares(unittest.TestCase):

    def armar_sae(self, filas=4):
        return pd.DataFrame({
            "Mes": [7] * filas,
            "Dia": [1] * filas,
            "Hora": list(range(1, filas + 1)),
            "Minutos": [0] * filas,
            "Hora Mes": list(range(1, filas + 1)),
            "Cuarto de Hora": list(range(1, filas + 1)),
            "clave": ["SAE UNO"] * filas,
            "intervalo": pd.to_datetime([
                f"2026-07-01 0{h}:00:00" for h in range(1, filas + 1)
            ]),
            "Gen_Unidad": [5.0, -3.0, 2.0, -1.0][:filas],
        })

    def armar_soc(self, filas=4):
        return pd.DataFrame({
            "central": ["SAE UNO"] * filas,
            "nombre_scada_original": ["SAE UNO"] * filas,
            "timestamp": pd.to_datetime([
                f"2026-07-01 0{h}:00:00" for h in range(1, filas + 1)
            ]),
            "soc": [0.5] * filas,
        })

    def test_la_hoja_sale_sin_vacias_sin_clave_auxiliar_y_sin_la_copia(self):
        df, _ = nucleo.construir_medidores(
            self.armar_sae(), self.armar_soc(), 7, registrar=lambda *a: None
        )

        self.assertEqual(
            list(df.columns), nucleo.COLUMNAS_MEDIDORES_SALIDA
        )

        for auxiliar in nucleo.COLUMNAS_AUXILIARES_MEDIDORES:
            self.assertNotIn(auxiliar, df.columns)

    def test_las_columnas_que_quedan_son_las_de_siempre_en_el_mismo_orden(self):
        """
        Quitar auxiliares no puede reordenar lo que queda: el orden
        sigue siendo el de LETRA_A_CAMPO (el de la planilla original).
        """

        quedan = [
            campo for campo in nucleo.LETRA_A_CAMPO.values()
            if campo not in nucleo.COLUMNAS_AUXILIARES_MEDIDORES
        ]

        self.assertEqual(nucleo.COLUMNAS_MEDIDORES_SALIDA, quedan)

    def test_la_copia_de_la_ventana_se_repone_al_leer(self):
        df, _ = nucleo.construir_medidores(
            self.armar_sae(), self.armar_soc(), 7, registrar=lambda *a: None
        )

        repuesto = nucleo.reponer_auxiliares_medidores(df)

        self.assertIn("Copia_Ventana", repuesto.columns)
        self.assertTrue(
            (repuesto["Copia_Ventana"] == repuesto["Ventana"]).all()
        )

    def test_un_libro_viejo_que_todavia_la_trae_se_respeta(self):
        """
        Consolidado_entradas.xlsx de una corrida anterior: la columna
        viene escrita y no hay que pisarla.
        """

        df = pd.DataFrame({
            "Ventana": [1, 1, 2],
            "Copia_Ventana": [9, 9, 9],
        })

        repuesto = nucleo.reponer_auxiliares_medidores(df)

        self.assertEqual(list(repuesto["Copia_Ventana"]), [9, 9, 9])


class TestColumnasDeSalidaDeLosCalculos(unittest.TestCase):

    def test_e_costos_sale_sin_el_ciclo_duplicado(self):
        self.assertNotIn("X", nucleo.COLUMNAS_SALIDA_E_COSTOS)

        # ...pero el dato no se pierde: sigue en la hoja una sola vez,
        # como "Ciclo de Carga del mes" (Copia_Ventana).
        nombres = [
            nucleo.NOMBRES_CALCULO_E_COSTOS[clave]
            for clave in nucleo.COLUMNAS_SALIDA_E_COSTOS
        ]
        self.assertEqual(nombres.count("Ciclo de Carga del mes"), 1)
        self.assertNotIn("Ciclo", nombres)

    def test_re545_sale_sin_la_columna_sin_nombre_ni_la_ventana_repetida(self):
        self.assertNotIn("BL", nucleo.COLUMNAS_SALIDA_RE545)
        self.assertNotIn("BR", nucleo.COLUMNAS_SALIDA_RE545)

        nombres = [
            nucleo.NOMBRES_CALCULO_RE545[clave]
            for clave in nucleo.COLUMNAS_SALIDA_RE545
        ]
        self.assertNotIn("", nombres)
        self.assertEqual(
            sum(n.lower() == "ventana de valorizacion" for n in nombres), 1
        )

    def test_los_intermedios_que_muestran_el_calculo_se_quedan(self):
        """
        El usuario pidio quitar lo trivial, no lo que deja ver de
        donde sale cada peso: las curvas monotonas, el ranking de CMg
        y las energias con FD siguen en la hoja.
        """

        for clave in ("W", "Y", "AB", "AC", "AD", "R", "AE", "AF", "AS", "AT"):
            self.assertIn(clave, nucleo.COLUMNAS_SALIDA_E_COSTOS)

        for clave in ("S", "U", "V", "BK", "BM", "BN", "BS", "BT"):
            self.assertIn(clave, nucleo.COLUMNAS_SALIDA_RE545)

    def test_los_grupos_siguen_cayendo_sobre_columnas_de_la_hoja(self):
        """
        Los encabezados de grupo se ubican con .index() sobre el orden
        de salida: si un grupo nombra una columna que ya no se
        escribe, escribir el libro revienta.
        """

        for grupos, columnas in (
            (nucleo.GRUPOS_CALCULO_E_COSTOS, nucleo.COLUMNAS_SALIDA_E_COSTOS),
            (nucleo.GRUPOS_CALCULO_RE545, nucleo.COLUMNAS_SALIDA_RE545),
        ):
            for etiqueta, claves in grupos:
                posiciones = [columnas.index(clave) for clave in claves]
                self.assertEqual(
                    posiciones,
                    list(range(min(posiciones), max(posiciones) + 1)),
                    f"el grupo '{etiqueta}' quedo partido en dos",
                )


class TestLibroDePagosOrdenado(unittest.TestCase):

    def armar_calculo(self):
        filas = 3
        base = {
            clave: list(range(1, filas + 1))
            for clave in nucleo.COLUMNAS_SALIDA_E_COSTOS
        }
        base["clave"] = ["SAE UNO"] * filas
        base["Barra"] = ["BARRA A"] * filas
        base["Energia_Positiva"] = [1.5, 2.5, 3.5]
        df = pd.DataFrame(base)
        return (
            df[nucleo.COLUMNAS_SALIDA_E_COSTOS]
            .rename(columns=nucleo.NOMBRES_CALCULO_E_COSTOS)
        )

    def escribir(self, carpeta):
        ruta = Path(carpeta) / "Pagos_BESS.xlsx"
        nucleo.escribir_pagos_bess(
            ruta,
            df_ecostos=self.armar_calculo(),
            df_resumen=pd.DataFrame({
                "NOMBRE": ["EMPRESA UNO"],
                "RECIBE": [1000.5],
                "PAGA": [0.0],
                "NETO": [1000.5],
            }),
            registrar=lambda *a: None,
        )
        return ruta

    def test_el_resumen_queda_primero_y_el_calculo_detras(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = self.escribir(carpeta)
            hojas = pd.ExcelFile(ruta).sheet_names

            self.assertEqual(hojas[0], nucleo.HOJA_RESUMEN)
            self.assertLess(
                hojas.index(nucleo.HOJA_RESUMEN),
                hojas.index(nucleo.HOJA_CALCULO_ECOSTOS),
            )

    def test_la_hoja_queda_legible_sin_tocar_los_datos(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = self.escribir(carpeta)
            wb = openpyxl.load_workbook(ruta)
            ws = wb[nucleo.HOJA_CALCULO_ECOSTOS]

            # nombres de columna en fila 2, en negrita y con el panel
            # inmovilizado justo debajo
            self.assertTrue(ws.cell(2, 1).font.bold)
            self.assertEqual(ws.freeze_panes, "A3")

            # ancho de columna puesto (pandas no pone ninguno)
            self.assertIsNotNone(ws.column_dimensions["A"].width)

            # y los datos son los mismos
            df = pd.read_excel(ruta, sheet_name=nucleo.HOJA_CALCULO_ECOSTOS,
                               header=1)
            self.assertEqual(list(df["Descarga kWh"]), [1.5, 2.5, 3.5])
            self.assertEqual(list(df["Configuracion"]), ["SAE UNO"] * 3)


class TestFormatoNoTocaLosDatos(unittest.TestCase):

    def test_formatear_hoja_deja_los_valores_como_estaban(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "x.xlsx"
            df = pd.DataFrame({
                "texto": ["a", "b"],
                "entero": [1, 2],
                "decimal": [1.25, 2.5],
                "vacia": [None, None],
            })
            df.to_excel(ruta, index=False)

            wb = openpyxl.load_workbook(ruta)
            ws = wb.active
            antes = [[c.value for c in fila] for fila in ws.iter_rows()]

            nucleo.formatear_hoja(ws)

            self.assertEqual(
                [[c.value for c in fila] for fila in ws.iter_rows()], antes
            )
            self.assertEqual(ws.cell(2, 2).number_format, "#,##0")
            self.assertEqual(ws.cell(2, 3).number_format, "#,##0.00")

    def test_una_hoja_vacia_no_rompe(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "v.xlsx"
            pd.DataFrame().to_excel(ruta, index=False)

            wb = openpyxl.load_workbook(ruta)
            nucleo.formatear_libro(wb)  # no debe levantar


if __name__ == "__main__":
    unittest.main()


class TestNoSeAbrenPlanillasDeMas(unittest.TestCase):
    """
    Pedido del usuario: "que no se abran planillas innecesarias" y
    "si alguna informacion esta en el consolidado que se saque de ahi".
    """

    def test_el_cmg_de_los_pagos_sale_del_consolidado(self):
        """
        La hoja 'CMg' del consolidado y cmg.xlsx son el mismo dato: se
        lee el consolidado (la foto de las entradas), no el archivo
        original.
        """

        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "Consolidado_entradas.xlsx"
            cmg = pd.DataFrame({
                "A": ["x"], "B": ["y"], "C": ["z"],
                "D_Barra": ["BARRA A"], "E": [""], "F_CMg": [50.0],
                "G": [""], "H_Cuarto": [1], "I_CMgProm": [45.0],
            })
            with pd.ExcelWriter(ruta) as w:
                cmg.to_excel(w, sheet_name="CMg", index=False)

            leido = nucleo.leer_cmg_consolidado(ruta, registrar=lambda *a: None)

            self.assertEqual(len(leido.columns), 9)
            self.assertEqual(
                nucleo.construir_dic_cmg(leido),
                nucleo.construir_dic_cmg(cmg),
            )

    def test_sin_la_hoja_cmg_el_error_dice_que_hay_que_generarla(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "Consolidado_entradas.xlsx"
            pd.DataFrame({"a": [1]}).to_excel(
                ruta, sheet_name="Medidores", index=False
            )

            with self.assertRaises(nucleo.ErrorEntrada) as fallo:
                nucleo.leer_cmg_consolidado(ruta, registrar=lambda *a: None)

            self.assertIn("CMg", str(fallo.exception))
